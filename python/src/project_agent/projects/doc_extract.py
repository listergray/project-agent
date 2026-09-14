"""???????Word / Excel / PDF ?? + ?????"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Sequence, Tuple

from loguru import logger

# (mime, bytes, label)
ImageBlob = Tuple[str, bytes, str]

_DOC_EXTS = {".md", ".markdown", ".txt", ".pdf", ".docx", ".xlsx", ".xlsm"}
_ZIP_DOC_EXTS = _DOC_EXTS | {".doc", ".xls"}  # zip ???????
_MAX_IMAGES = 8
_MAX_IMAGE_BYTES = 8 * 1024 * 1024
_IMG_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".emf", ".wmf", ".tif", ".tiff"}


@dataclass
class ExtractResult:
    text: str
    images: List[ImageBlob] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)


def supported_doc_exts() -> set[str]:
    return set(_DOC_EXTS)


def is_supported_doc(name_or_path: str | Path) -> bool:
    return Path(name_or_path).suffix.lower() in _DOC_EXTS


def _guess_mime(name: str, raw: bytes) -> str:
    ext = Path(name).suffix.lower()
    if ext in {".jpg", ".jpeg"}:
        return "image/jpeg"
    if ext == ".gif":
        return "image/gif"
    if ext == ".webp":
        return "image/webp"
    if ext in {".bmp"}:
        return "image/bmp"
    if raw[:8].startswith(b"\x89PNG"):
        return "image/png"
    if raw[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if raw[:4] == b"GIF8":
        return "image/gif"
    return "image/png"


def _add_image(
    out: List[ImageBlob],
    seen: set[str],
    name: str,
    raw: bytes,
    *,
    label: str = "",
) -> None:
    if not raw or len(raw) > _MAX_IMAGE_BYTES or len(out) >= _MAX_IMAGES:
        return
    ext = Path(name).suffix.lower()
    # EMF/WMF ?? OCR ???????
    if ext in {".emf", ".wmf"}:
        return
    if ext and ext not in _IMG_EXTS and not (
        raw.startswith(b"\x89PNG") or raw[:3] == b"\xff\xd8\xff" or raw[:4] == b"GIF8"
    ):
        return
    key = f"{len(raw)}:{raw[:64].hex()}"
    if key in seen:
        return
    seen.add(key)
    mime = _guess_mime(name, raw)
    out.append((mime, raw, label or Path(name).name))


def _extract_md_txt(path: Path) -> ExtractResult:
    return ExtractResult(text=path.read_text(encoding="utf-8", errors="ignore"))


def _extract_docx(path: Path) -> ExtractResult:
    try:
        from docx import Document
    except ImportError as e:  # noqa: BLE001
        raise ValueError("?? Word ?? python-docx?????????") from e

    warnings: List[str] = []
    images: List[ImageBlob] = []
    seen: set[str] = set()
    lines: List[str] = []

    doc = Document(str(path))
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            lines.append(t)

    for ti, table in enumerate(doc.tables, start=1):
        lines.append(f"\n## ??{ti}")
        for row in table.rows:
            cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    # ????? zip ? word/media?? related_parts ???
    try:
        with zipfile.ZipFile(path, "r") as zf:
            media = [n for n in zf.namelist() if n.startswith("word/media/")]
            media.sort()
            for i, name in enumerate(media, start=1):
                try:
                    raw = zf.read(name)
                except Exception:  # noqa: BLE001
                    continue
                _add_image(images, seen, name, raw, label=f"Word?{i}:{Path(name).name}")
    except Exception as e:  # noqa: BLE001
        warnings.append(f"Word ????????: {e}")

    text = "\n".join(lines).strip()
    if not text and not images:
        warnings.append("Word ??????????")
    logger.info(f"[DocExtract] docx paras/tables chars={len(text)} images={len(images)}")
    return ExtractResult(text=text, images=images, warnings=warnings)


def _extract_xlsx(path: Path) -> ExtractResult:
    try:
        import openpyxl
    except ImportError as e:  # noqa: BLE001
        raise ValueError("?? Excel ?? openpyxl?????????") from e

    warnings: List[str] = []
    images: List[ImageBlob] = []
    seen: set[str] = set()
    parts: List[str] = []

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=False)
    try:
        for sheet in wb.worksheets:
            parts.append(f"## ???: {sheet.title}")
            rows_out: List[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = []
                empty = True
                for v in row:
                    if v is None:
                        cells.append("")
                    else:
                        empty = False
                        cells.append(str(v).strip())
                if not empty:
                    rows_out.append(" | ".join(cells))
            if rows_out:
                parts.extend(rows_out[:500])  # ??????????
                if len(rows_out) > 500:
                    parts.append(f"?????? {len(rows_out)} ??")
            else:
                parts.append("????")

            # openpyxl ???
            for i, img in enumerate(getattr(sheet, "_images", []) or [], start=1):
                try:
                    ref = getattr(img, "ref", None) or getattr(img, "path", None)
                    raw = None
                    if hasattr(img, "_data") and callable(img._data):
                        raw = img._data()
                    elif ref and hasattr(wb, "_archive"):
                        pass
                    if raw is None and hasattr(img, "image"):
                        # Pillow Image
                        buf = io.BytesIO()
                        img.image.save(buf, format="PNG")
                        raw = buf.getvalue()
                    if raw:
                        _add_image(
                            images,
                            seen,
                            f"{sheet.title}_img{i}.png",
                            raw,
                            label=f"Excel?:{sheet.title}-{i}",
                        )
                except Exception as e:  # noqa: BLE001
                    logger.debug(f"[DocExtract] xlsx sheet image skip: {e}")
    finally:
        wb.close()

    # ???xlsx ?? zip?? xl/media
    try:
        with zipfile.ZipFile(path, "r") as zf:
            media = [n for n in zf.namelist() if "/media/" in n.lower()]
            media.sort()
            for i, name in enumerate(media, start=1):
                try:
                    raw = zf.read(name)
                except Exception:  # noqa: BLE001
                    continue
                _add_image(images, seen, name, raw, label=f"Excel??{i}:{Path(name).name}")
    except Exception as e:  # noqa: BLE001
        warnings.append(f"Excel ????????: {e}")

    text = "\n".join(parts).strip()
    logger.info(f"[DocExtract] xlsx chars={len(text)} images={len(images)}")
    return ExtractResult(text=text, images=images, warnings=warnings)


def _extract_pdf(path: Path) -> ExtractResult:
    try:
        from pypdf import PdfReader
        import markdownify
    except ImportError as e:  # noqa: BLE001
        raise ValueError("?? PDF ?? pypdf") from e

    warnings: List[str] = []
    images: List[ImageBlob] = []
    seen: set[str] = set()
    pages_html: List[str] = []

    reader = PdfReader(str(path))
    for i, page in enumerate(reader.pages):
        txt = page.extract_text() or ""
        pages_html.append(f"<h1>?{i+1}?</h1>\n<p>{txt}</p>")
        # pypdf ???
        try:
            for j, img in enumerate(getattr(page, "images", []) or [], start=1):
                raw = getattr(img, "data", None)
                name = getattr(img, "name", None) or f"page{i+1}_{j}.png"
                if raw:
                    _add_image(
                        images,
                        seen,
                        str(name),
                        raw,
                        label=f"PDF?{i+1}??{j}",
                    )
        except Exception as e:  # noqa: BLE001
            logger.debug(f"[DocExtract] pdf page {i+1} images: {e}")

    md = markdownify.markdownify("\n".join(pages_html), heading_style="ATX")
    text = (md or "").strip()
    # ????????????? OCR??????????
    if len(re.sub(r"\s+", "", text)) < 20 and not images:
        warnings.append(
            "PDF ??????????????????????"
            "???????????????????????????? PDF"
        )
    elif len(re.sub(r"\s+", "", text)) < 20 and images:
        warnings.append("PDF ??????????????? OCR")

    logger.info(f"[DocExtract] pdf pages={len(reader.pages)} chars={len(text)} images={len(images)}")
    return ExtractResult(text=text, images=images, warnings=warnings)


def extract_document(path: Path) -> ExtractResult:
    """??????????????"""
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        return _extract_md_txt(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _extract_xlsx(path)
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in {".doc", ".xls"}:
        raise ValueError(
            f"?????? {suffix}????? .docx / .xlsx ????"
        )
    raise ValueError(f"?????????: {suffix}")


def extract_document_text(path: Path) -> str:
    """????? RAG / ???????"""
    return extract_document(path).text


def zip_preferred_docs() -> Sequence[str]:
    # ????md > docx > xlsx > pdf > txt
    order = [
        ".md",
        ".markdown",
        ".docx",
        ".xlsx",
        ".xlsm",
        ".pdf",
        ".txt",
    ]
    return order
